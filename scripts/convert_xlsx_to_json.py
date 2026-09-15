#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把「瑪奇物資.xlsx」轉換成網頁模擬器使用的 data.json
用法: python3 scripts/convert_xlsx_to_json.py data.xlsx data.json
"""
import sys
import json
import re
import datetime
import openpyxl

def build_merge_map(ws):
    merge_map = {}
    for mrange in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = mrange.bounds
        top_val = ws.cell(row=min_row, column=min_col).value
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                merge_map[(r, c)] = top_val
    return merge_map

def cell_value(ws, merge_map, row, col):
    if (row, col) in merge_map:
        return merge_map[(row, col)]
    return ws.cell(row=row, column=col).value

def parse_resource(text):
    """'原木*10' -> {'name': '原木', 'qty': 10}"""
    if text is None:
        return None
    text = str(text).strip()
    if not text:
        return None
    m = re.match(r'^(.*?)[\*×xX](\d+)\s*$', text)
    if m:
        name = m.group(1).strip()
        qty = int(m.group(2))
        return {"name": name, "qty": qty}
    return {"name": text, "qty": 1}

def parse_special_shop(text):
    if text is None:
        return None
    text = str(text).strip()
    if not text:
        return None
    return text.strip('()（）')

def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "data.xlsx"
    dst = sys.argv[2] if len(sys.argv) > 2 else "data.json"

    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb.worksheets[0]
    merge_map = build_merge_map(ws)

    max_row = ws.max_row
    max_col = ws.max_column

    # 找到欄位標題列 (找出含有「地區」字樣的那一列)
    header_row = None
    col_map = {}
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            v = cell_value(ws, merge_map, r, c)
            if v == "地區":
                header_row = r
                break
        if header_row:
            break

    if header_row is None:
        raise RuntimeError("找不到標題列 (地區/NPC/兌換物品/所需資源(最大值)/特殊商店)")

    # 標題列可能欄位名稱有些微差異，用包含關係比對
    header_aliases = {
        "地區": "region",
        "NPC": "npc",
        "兌換物品": "item",
        "所需資源": "resource",   # 涵蓋「所需資源(最大值)」等變體
        "特殊商店": "shop",
    }
    for c in range(1, max_col + 1):
        v = cell_value(ws, merge_map, header_row, c)
        if not v:
            continue
        v = str(v)
        for key, mapped in header_aliases.items():
            if key in v:
                col_map[mapped] = c
                break

    required = ["region", "npc", "item", "resource"]
    for req in required:
        if req not in col_map:
            raise RuntimeError(f"標題列缺少欄位: {req}")

    rows = []
    for r in range(header_row + 1, max_row + 1):
        region = cell_value(ws, merge_map, r, col_map["region"])
        npc = cell_value(ws, merge_map, r, col_map["npc"])
        item = cell_value(ws, merge_map, r, col_map["item"])
        resource_raw = ws.cell(row=r, column=col_map["resource"]).value  # 不用 merge_map，每列各自的花費
        shop_raw = ws.cell(row=r, column=col_map.get("shop", 0)).value if col_map.get("shop") else None

        if region is None and npc is None and item is None and resource_raw is None and shop_raw is None:
            continue

        rows.append({
            "region": str(region).strip() if region else None,
            "npc": str(npc).strip() if npc else None,
            "item": str(item).strip() if item else None,
            "resource": parse_resource(resource_raw),
            "shop": parse_special_shop(shop_raw),
        })

    items = {}
    notes = []

    for row in rows:
        if row["item"] and row["resource"]:
            entry = {
                "region": row["region"],
                "npc": row["npc"],
                "resource": row["resource"],
            }
            if row["shop"]:
                entry["shop"] = row["shop"]
            items.setdefault(row["item"], []).append(entry)
        elif row["shop"] and not row["item"]:
            notes.append({
                "region": row["region"],
                "npc": row["npc"],
                "shop": row["shop"],
            })

    # 去除完全重複的 recipe
    for k, v in items.items():
        seen = set()
        deduped = []
        for e in v:
            key = json.dumps(e, ensure_ascii=False, sort_keys=True)
            if key not in seen:
                seen.add(key)
                deduped.append(e)
        items[k] = deduped

    output = {
        "generatedAt": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "items": items,
        "notes": notes,
    }

    with open(dst, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    print(f"轉換完成：{len(items)} 種兌換物品，{len(notes)} 筆額外備註 -> {dst}")

if __name__ == "__main__":
    main()
